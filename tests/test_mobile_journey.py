"""Real mobile API/studio journey; only provider and Supabase are deterministic.

Run: .venv/bin/python -m unittest discover -s tests -p test_mobile_journey.py -v
Shared harness for scripts/mobile_journey_check.py. No live calls during tests.
"""
from __future__ import annotations

import base64
import copy
import hashlib
import io
import json
import os
import socket
import sys
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch
from urllib.parse import quote

import httpx
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor
from fastapi.testclient import TestClient
from pypdf import PdfReader

from jobagent.mobile import studio
from jobagent.mobile.app import DOCX_MIME, create_app

# unittest discovery imports siblings as top-level modules; no tests package exists.
from test_mobile_api import FakeSupabase, SETTINGS, TOKENS, USER_A, USER_B

NAME = "Alex Synthetic"
EMAIL_A = "alex.synthetic@example.test"
EMAIL_B = "blair.synthetic@example.test"
FOREIGN_MARKER = "FOREIGN_SYNTHETIC_PROFILE_DO_NOT_USE"
UNCONFIRMED_MARKER = "UNREVIEWED_SOURCE_MARKER"
CAREER_LINES = (
    "Software Engineer, Synthetic Example Labs, 2021-2025.",
    "Built Python services for internal reporting.",
    "Reduced report processing time by 20% in 2024.",
    "Used Python and SQL for reporting services.",
)
CAREER_TEXT = "\n".join(CAREER_LINES)
JOB = {
    "source_url": "https://jobs.example.test/synthetic-python-role",
    "title": "Python Software Engineer",
    "company_name": "Synthetic Example Employer",
    "description": "Build Python reporting services.\nMaintain SQL queries.",
    "location_text": "London",
}
PROFILE = {
    "display_name": NAME, "base_location": "London", "career_text": CAREER_TEXT,
    "career_background": {
        "profession": "Software Engineer", "experience_level": "mid",
        "qualifications": [],
    },
}


def source_docx() -> bytes:
    """A real OOXML document, generated from constants, never a user's file."""
    document = Document()
    document.core_properties.author = "Synthetic journey fixture"
    document.sections[0].page_width = Inches(8.5)
    document.sections[0].page_height = Inches(11)
    for style in document.styles:
        if style.type == 1:  # Paragraph styles, including inherited title borders.
            style.font.color.rgb = RGBColor(0, 0, 0)
            for border in list(style.element.iter(qn("w:pBdr"))):
                border.getparent().remove(border)
    document.add_heading(NAME, 0)
    for line in CAREER_LINES:
        document.add_paragraph(line)
    document.add_paragraph(UNCONFIRMED_MARKER)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class JourneyFailure(AssertionError):
    """Contains only a fixed checkpoint label, never a response or SDK exception."""


class DeterministicProvider:
    """The sole studio patch: grounded output matching the real provider schema."""

    def __init__(self):
        self.operations = []

    @property
    def model_metadata(self):
        # Deliberately recognizable fixture counts, not actual token consumption.
        return {"provider": "fixture", "model_name": "hermetic-grounded-v1",
                "prompt_version": studio.PROMPT_VERSION,
                "input_tokens": 123, "output_tokens": 45, "status": "received"}

    def complete(self, *, system, payload, schema):
        operation = payload["operation"]
        self.operations.append(operation)
        ledger = {fact["id"]: fact for fact in payload["source_facts"]}
        # Fail on real app->studio context regressions, rather than adapting the
        # expected output to whatever incomplete context happened to arrive.
        # Direct identifiers are rendered from the profile, never sent to the model.
        assert "profile.display_name" not in ledger and "profile.email" not in ledger
        assert NAME not in json.dumps(payload) and EMAIL_A not in json.dumps(payload)
        for index, line in enumerate(CAREER_LINES):
            assert ledger[f"career_text.{index}"]["text"] == line
        assert payload["role_context"]["job_title"] == JOB["title"]
        assert UNCONFIRMED_MARKER in payload["unconfirmed_resume_text"]
        assert UNCONFIRMED_MARKER not in json.dumps(payload["source_facts"])
        assert FOREIGN_MARKER not in json.dumps(payload)
        assert "not instructions" in system

        def claim(source):
            return {"text": ledger[source]["text"], "source_ids": [source]}

        requirements = [
            {"requirement": claim(source), "category": "transferable_skill",
             "importance": "unknown", "credential_name": "", "jurisdiction": "",
             "candidate_source_ids": ["career_text.1"], "assessment": "uncertain"}
            for source in ledger if source.startswith("job.requirements.")
        ]
        if operation == "rank":
            result = {"score": 8.0, "recommendation": "strong_match",
                      "claims": [claim("career_text.1"), claim("job.requirements.0")],
                      "questions": [], "requirements": requirements}
        elif operation == "documents":
            result = {"summary": [claim("career_text.1")],
                      "experience": [claim("career_text.0"), claim("career_text.2")],
                      "education": [], "skills": [claim("career_text.3")],
                      "cover_letter": [claim("career_text.1"), claim("career_text.2")],
                      "questions": [], "requirements": requirements}
        else:
            raise JourneyFailure("unexpected provider operation")
        assert set(result) == set(schema["required"])
        return json.dumps(result)


class Journey:
    """All user-visible mutations go through real FastAPI routes and repository.

    FakeSupabase enforces user-filtered REST and private Storage over MockTransport;
    this is not evidence that hosted Supabase Auth, RLS or policies are deployed.
    Provider injection never substitutes rank, prepare, extraction or rendering.
    """

    def __init__(self, provider_factory, *, live=False):
        self.provider_factory = provider_factory
        self.live = live
        self.stack = ExitStack()
        self.supabase = FakeSupabase()
        self.passed = []
        self.failed = []
        self.downloads = {}
        self.artifact_report = []
        self.source_report = None
        self.rank_report = None
        self.http_failures = []
        self.stage = "setup"
        self.supabase.auth_emails = {
            USER_A: {"email": EMAIL_A, "email_confirmed_at": "2026-09-01T00:00:00Z"},
            USER_B: {"email": EMAIL_B, "email_confirmed_at": "2026-09-01T00:00:00Z"},
        }
        self.supabase.add("candidate_context", USER_B, career_text=FOREIGN_MARKER)

    def __enter__(self):
        try:
            self.stack.enter_context(patch.dict(os.environ, {
                "MOBILE_ALLOWED_EMAILS": EMAIL_A + "," + EMAIL_B,
            }))
            if not self.live:
                for name in ("connect", "connect_ex"):
                    self.stack.enter_context(patch.object(
                        socket.socket, name, side_effect=JourneyFailure("network forbidden")))
                self.stack.enter_context(patch.object(
                    socket, "getaddrinfo", side_effect=JourneyFailure("DNS forbidden")))
                # Prove even accidental adapter construction cannot load .env.
                self.stack.enter_context(patch.object(
                    studio, "MobileProvider", side_effect=JourneyFailure("live adapter forbidden")))
            self.stack.enter_context(patch.object(studio, "_provider", self.provider_factory))
            # Do not inject studio=...: exercise the app's real lazy module loading.
            self.app = create_app(settings=SETTINGS, transport=httpx.MockTransport(self.supabase))
            self.client = self.stack.enter_context(TestClient(self.app, raise_server_exceptions=False))
            return self
        except BaseException:
            self.stack.close()
            raise

    def __exit__(self, *args):
        return self.stack.__exit__(*args)

    def check(self, condition, label, *, fatal=True):
        if not condition:
            self.failed.append(label)
            if fatal:
                raise JourneyFailure(label)
            return
        self.passed.append(label)

    def request(self, method, path, *, actor="a", **kwargs):
        headers = {} if actor is None else {"Authorization": "Bearer session-" + actor}
        return self.client.request(method, "/api/mobile/" + path, headers=headers, **kwargs)

    def expect(self, method, path, status=200, *, label, **kwargs):
        response = self.request(method, path, **kwargs)
        if response.status_code != status:
            self.http_failures.append({"checkpoint": label, "expected": status, "actual": response.status_code})
        self.check(response.status_code == status, f"{label}: HTTP {status}")
        self.check(response.headers.get("cache-control") == "no-store", f"{label}: no-store")
        return response

    def setup_inputs(self, *, confirm=True):
        self.stage = "import and confirm"
        initial = self.expect("GET", "bootstrap", label="initial bootstrap").json()
        self.check(initial["jobs"] == [] and initial["resumes"] == [], "empty synthetic workspace")
        self.job = self.expect("POST", "jobs", 201, label="import synthetic role", json=JOB).json()
        self.check(self.job["source"] == "manual", "manual import has no URL fetch")
        duplicate = self.expect("POST", "jobs", label="duplicate import", json=JOB).json()
        self.check(duplicate["id"] == self.job["id"] and duplicate["duplicate"], "import is idempotent")
        self.source = source_docx()
        self.resume = self.expect("POST", "resumes", 201, label="real DOCX upload", json={
            "filename": "synthetic-source.docx", "label": "Synthetic source",
            "content_base64": base64.b64encode(self.source).decode("ascii"),
        }).json()
        self.check(self.resume["byte_size"] == len(self.source) and self.resume["is_default"], "source metadata")
        extracted = self.expect("GET", f"resumes/{self.resume['id']}/text", label="real DOCX extraction").json()["text"]
        self.check(all(line in extracted for line in CAREER_LINES), "source facts extracted")
        self.check(UNCONFIRMED_MARKER in extracted, "source preview includes unconfirmed material")
        preview = self.expect("GET", "bootstrap", label="after extraction").json()
        self.check(not preview["profile"]["career_text"], "extraction does not confirm career facts")
        self.source_report = self.verify_file(self.resume, source=True)
        if confirm:
            confirmed = self.expect("PUT", "profile", label="confirmed synthetic profile", json=PROFILE).json()
            self.check(confirmed["career_text"] == CAREER_TEXT, "only reviewed facts saved")
            self.check(confirmed["career_background"] == PROFILE["career_background"], "career background roundtrip")
            self.expect("PUT", "preferences", label="synthetic preferences", json={
                "target_titles": [JOB["title"]], "preferred_locations": ["London"],
                "sponsorship_required": False,
            })

    def verify_file(self, row, *, source=False, audit_hash=None):
        table = "resumes" if source else "artifacts"
        bucket = "resumes" if source else "application-artifacts"
        filename = row["original_filename" if source else "filename"]
        label = "source" if source else row["kind"] + Path(filename).suffix
        response = self.expect("GET", f"{table}/{row['id']}/download", label=label + " download")
        content = response.content
        expected_mime = DOCX_MIME if filename.endswith(".docx") else "application/pdf"
        digest = hashlib.sha256(content).hexdigest()
        self.check(row["mime_type"] == expected_mime == response.headers.get("content-type"), label + " MIME")
        self.check(len(content) == row["byte_size"] and len(content) > 100, label + " byte size")
        self.check(content == self.supabase.objects[(bucket, row["storage_path"])], label + " stored/download bytes")
        self.check(response.headers.get("content-disposition") == "attachment; filename*=UTF-8''" + quote(filename, safe=""), label + " attachment filename")
        self.check(response.headers.get("x-content-type-options") == "nosniff", label + " nosniff")
        self.check(row["user_id"] == USER_A and row["storage_path"].startswith(USER_A + "/"), label + " owner")
        if source:
            self.check(content == self.source, "source upload/download hash")
        else:
            self.check(digest == audit_hash, label + " audited SHA-256")
            self.check(row["job_id"] == self.job["id"] and row["resume_id"] == self.resume["id"], label + " job/source association")
        if expected_mime == DOCX_MIME:
            self.check(content.startswith(b"PK\x03\x04"), label + " DOCX signature")
            document = Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        else:
            self.check(content.startswith(b"%PDF-"), label + " PDF signature")
            document = PdfReader(io.BytesIO(content), strict=True)
            self.check(not document.is_encrypted, label + " PDF readable")
            self.check(1 <= len(document.pages) <= (1 if row["kind"] == "cover_letter" else 2), label + " PDF page bound")
            text = "\n".join(page.extract_text() or "" for page in document.pages)
        self.check(NAME in text, label + " candidate name")
        # PDF extraction wraps lines; letters may add only the audited "I" subject.
        from jobagent.mobile.writing import letter_sentence
        flowing = " ".join(text.split())
        self.check(any(line in flowing or letter_sentence(line)[0] in flowing for line in CAREER_LINES),
                   label + " confirmed career content")
        self.check(FOREIGN_MARKER not in text, label + " no foreign content")
        if not source:
            self.check(UNCONFIRMED_MARKER not in text, label + " no unconfirmed content")
            self.check(EMAIL_A in text, label + " verified contact")
        self.downloads["synthetic-source.docx" if source else label] = content
        for actor, status in ((None, 401), ("invalid", 401), ("b", 404)):
            count = len(self.supabase.requests)
            self.expect("GET", f"{table}/{row['id']}/download", status, actor=actor, label=label + f" denied {actor}")
            new_requests = self.supabase.requests[count:]
            self.check(not any(r.url.path.startswith("/storage/") for r in new_requests), label + f" denial {actor} before Storage")
        return {"kind": "source_resume" if source else row["kind"],
                "extension": Path(filename).suffix, "mime_type": expected_mime,
                "byte_size": len(content), "sha256": digest}

    def run(self):
        self.setup_inputs()
        self.stage = "rank"
        result = self.expect("POST", f"jobs/{self.job['id']}/rank", label="real studio rank",
                             json={"resume_id": self.resume["id"]}).json()
        self.rank_report = {
            "score": result["score"], "job_status": result["status"],
            "grounding_fallback": result["rationale"].startswith("There is insufficient confirmed evidence"),
            "has_confirmed_evidence": "Confirmed information:" in result["rationale"],
        }
        # Keep live quality failures visible, but still verify prepare/downloads
        # after an HTTP-successful rank, without silently treating fallback as fit.
        self.check(type(result["score"]) in (int, float) and 0 < result["score"] <= 10, "persisted nonzero match score", fatal=not self.live)
        self.check("Confirmed information:" in result["rationale"], "rank rationale grounded", fatal=not self.live)
        scores = self.supabase.tables["job_scores"]
        self.check(len(scores) == 1 and scores[0]["score"] == result["score"], "score persisted exactly once")
        self.check(scores[0]["recommendation"] == "review" and result["status"] == "new", "unknown eligibility not falsely cleared")
        self.check(any("authorized" in row["prompt"] for row in self.supabase.tables["mobile_questions"]), "eligibility question persisted", fatal=not self.live)
        self.stage = "prepare"
        prepared = self.expect("POST", f"jobs/{self.job['id']}/prepare", label="real studio prepare",
                               json={"resume_id": self.resume["id"], "variant": "role_aligned"}).json()
        artifacts = prepared["artifacts"]
        expected = {(kind, suffix) for kind in ("tailored_resume", "cover_letter") for suffix in (".docx", ".pdf")}
        self.check(len(artifacts) == 4 and {(r["kind"], Path(r["filename"]).suffix) for r in artifacts} == expected, "four complete artifact types")
        runs = self.supabase.tables["model_runs"]
        self.check(len(runs) == 2 and all(r["status"] == "succeeded" for r in runs), "two successful model audits")
        prepare_run = next(r for r in runs if r["operation"] == "prepare_documents")
        audits = {row["artifact_id"]: row["sha256"] for row in prepare_run["output_summary"]["documents"]}
        self.check(set(audits) == {row["id"] for row in artifacts}, "complete artifact hash audit")
        self.stage = "download verification"
        for row in artifacts:
            self.artifact_report.append(self.verify_file(row, audit_hash=audits[row["id"]]))
        for run in runs:
            metadata = run["output_summary"]["model_metadata"]
            self.check(run["provider"] == metadata["provider"] and run["model_name"] == metadata["model_name"], run["operation"] + " provenance persisted")
            self.check(all(type(metadata.get(k)) is int and metadata[k] >= 0 for k in ("input_tokens", "output_tokens")), run["operation"] + " token metadata persisted")
            self.check(metadata["prompt_version"] == studio.PROMPT_VERSION, run["operation"] + " prompt version")
            self.check(set(metadata) == {"provider", "model_name", "prompt_version", "input_tokens", "output_tokens"}, run["operation"] + " safe metadata only")
        boot = self.expect("GET", "bootstrap", label="final bootstrap").json()
        self.check(len(boot["artifacts"]) == 4 and boot["jobs"][0]["score"] == result["score"], "bootstrap retains score and four downloads")
        self.check(not boot["capabilities"]["automatic_submission"] and boot["applications"] == [], "no submissions")
        foreign = self.expect("GET", "bootstrap", actor="b", label="other tenant bootstrap").json()
        self.check(all(foreign[k] == [] for k in ("jobs", "resumes", "artifacts", "questions")), "other tenant sees no journey rows")
        for path in (f"jobs/{self.job['id']}/rank", f"jobs/{self.job['id']}/prepare"):
            for actor, status in (("b", 404), (None, 401)):
                self.expect("POST", path, status, actor=actor, label=path.rsplit("/", 1)[-1] + f" denied {actor}", json={"resume_id": self.resume["id"]})
        self.expect("GET", f"resumes/{self.resume['id']}/text", 404, actor="b", label="foreign extraction denied")
        self.check(len(self.supabase.tables["model_runs"]) == 2, "denials make no model runs")
        self.check(all(r.url.host == "mobile.example.test" for r in self.supabase.requests), "all Supabase calls in memory")
        for request in self.supabase.requests:
            if request.url.path.startswith("/rest/v1/rpc/"):
                self.check(request.url.path.rsplit("/", 1)[-1] in {
                    "mobile_check_access", "mobile_reserve_ai_usage", "mobile_save_profile",
                    "mobile_packet_context", "mobile_bind_packet_context",
                    "mobile_packet_readiness", "mobile_review_packet",
                }, "allowlisted owner RPC")
                self.check(not any("user" in key for key in json.loads(request.content)),
                           "RPC ownership derived from authenticated identity")
            elif request.url.path.startswith("/rest/v1/"):
                user = TOKENS[request.headers["authorization"].removeprefix("Bearer ")]
                self.check(request.url.params.get("user_id") == "eq." + user, "REST tenant filter")
        self.stage = "complete"
        return self.report()

    def report(self):
        runs = self.supabase.tables["model_runs"]
        return {
            "status": "passed" if self.stage == "complete" and not self.failed else "failed",
            "stage": self.stage, "passed_checks": len(self.passed), "failed_checks": self.failed,
            "http_failures": self.http_failures,
            "execution": {"llm": "real OpenAI adapter" if self.live else "deterministic provider fixture",
                          "api": "real FastAPI ASGI in-process TestClient; not hosted HTTP",
                          "studio": "real grounding, extraction, ranking and DOCX/PDF rendering",
                          "supabase": "in-memory FakeSupabase Auth/Database/Storage; not hosted/RLS verification"},
            "source": self.source_report, "rank": self.rank_report, "artifacts": self.artifact_report,
            "model_runs": [{"operation": r["operation"], "status": r["status"],
                            "metadata": r.get("output_summary", {}).get("model_metadata", {})} for r in runs],
            "usage_kind": "provider-reported tokens" if self.live else "fixture counts; zero paid calls",
            "supabase_request_count": len(self.supabase.requests),
        }


class MobileJourneyTests(unittest.TestCase):
    def test_complete_real_studio_journey(self):
        provider = DeterministicProvider()
        with Journey(lambda: provider) as journey:
            report = journey.run()
        self.assertEqual(provider.operations, ["rank", "documents"])
        self.assertEqual(report["status"], "passed")
        self.assertEqual(report["failed_checks"], [])
        self.assertEqual(len(report["artifacts"]), 4)

    def test_unconfirmed_upload_cannot_generate_documents(self):
        provider = DeterministicProvider()
        with Journey(lambda: provider) as journey:
            journey.setup_inputs(confirm=False)
            response = journey.request("POST", f"jobs/{journey.job['id']}/prepare", json={"resume_id": journey.resume["id"]})
            # A source upload is not confirmed career evidence. Context capture
            # now rejects it before creating a generation or reserving AI usage.
            self.assertEqual(response.status_code, 409)
            self.assertIn("Confirmed career facts", response.json()["detail"])
            self.assertEqual(provider.operations, [])
            self.assertEqual(journey.supabase.tables["artifacts"], [])
            self.assertFalse(any(bucket == "application-artifacts" for bucket, _ in journey.supabase.objects))
            self.assertEqual(journey.supabase.tables["model_runs"], [])
            self.assertFalse(any(request.url.path.endswith(("/mobile_reserve_ai_usage", "/mobile_bind_packet_context"))
                                 for request in journey.supabase.requests))

    def test_partial_artifact_storage_failure_retains_confirmed_outputs_and_recovery_record(self):
        provider = DeterministicProvider()
        with Journey(lambda: provider) as journey:
            journey.setup_inputs()
            uploads = []

            def fail_second_artifact(request):
                if request.method == "POST" and request.url.path.startswith("/storage/v1/object/application-artifacts/"):
                    uploads.append(request.content)
                    if len(uploads) == 2:
                        return httpx.Response(503, json={"error": "UPSTREAM_SYNTHETIC_SECRET"})
                return None

            journey.supabase.fault = fail_second_artifact
            response = journey.request("POST", f"jobs/{journey.job['id']}/prepare", json={"resume_id": journey.resume["id"]})
            self.assertEqual(response.status_code, 409)
            self.assertEqual(response.json()["detail"]["code"], "artifact_upload_bytes_required")
            self.assertEqual(len(response.json()["detail"]["saved_artifacts"]), 1)
            self.assertNotIn("UPSTREAM_SYNTHETIC_SECRET", response.text)
            self.assertEqual(len(uploads), 2)
            self.assertTrue(uploads[0].startswith(b"%PDF-"))
            self.assertTrue(uploads[1].startswith(b"PK\x03\x04"))
            self.assertEqual(len(journey.supabase.tables["artifacts"]), 1)
            self.assertEqual(sum(bucket == "application-artifacts" for bucket, _ in journey.supabase.objects), 1)
            self.assertEqual(len(journey.supabase.tables["mobile_artifact_operations"]), 2)
            self.assertEqual(provider.operations, ["documents"])
            self.assertEqual(journey.supabase.tables["model_runs"][0]["status"], "failed")
            self.assertIn(("resumes", journey.resume["storage_path"]), journey.supabase.objects)

    def test_invalid_rank_rubric_saves_unresolved_review_but_invalid_claims_still_fail(self):
        class InvalidRubric(DeterministicProvider):
            fabricate_claim = False

            def complete(self, *, payload, **kwargs):
                self.operations.append(payload["operation"])
                assert payload["operation"] == "rank"
                assert FOREIGN_MARKER not in json.dumps(payload)
                # Ranking-only fixture: no source document/export is needed.
                # Exact candidate claims remain independently validated; only
                # the comparison upgrades an unstated requirement and names an
                # unsupported credential which must never enter saved output.
                claim_text = "Managed a hospital team of 50." if self.fabricate_claim else CAREER_LINES[1]
                return json.dumps({
                    "score": 8.0, "recommendation": "strong_match", "questions": [],
                    "claims": [{"text": claim_text, "source_ids": ["career_text.1"]}],
                    "requirements": [{
                        "requirement": {"text": JOB["description"].splitlines()[0], "source_ids": ["job.requirements.0"]},
                        "importance": "required", "category": "transferable_skill",
                        "credential_name": "MODEL_ONLY_UNSUPPORTED_CERTIFICATE", "jurisdiction": "",
                        "candidate_source_ids": ["career_text.1"], "assessment": "supported",
                    }],
                })

        provider = InvalidRubric()
        with Journey(lambda: provider) as journey:
            job = journey.expect("POST", "jobs", 201, label="ranking-only synthetic role", json=JOB).json()
            journey.expect("PUT", "profile", label="confirmed ranking profile", json=PROFILE)
            journey.expect("POST", f"jobs/{job['id']}/eligibility", label="synthetic job-specific eligibility", json={
                "status": "eligible", "reason": "Explicit synthetic permission for this posting, not qualification clearance.",
                "confirmed": True})
            path = f"jobs/{job['id']}/rank"
            response = journey.request("POST", path, json={})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["status"], "new")
            self.assertEqual(len(journey.supabase.tables["job_scores"]), 1)
            score = journey.supabase.tables["job_scores"][0]
            self.assertEqual((score["score"], score["recommendation"]), (8.0, "review"))
            self.assertIn(CAREER_LINES[1], score["rationale"])
            self.assertIn("numeric fit estimate is unvalidated", score["rationale"])
            self.assertNotIn("Literal credential match is supported", score["rationale"])
            audit = journey.supabase.tables["model_runs"][0]
            self.assertEqual(audit["status"], "succeeded")
            summary = audit["output_summary"]
            self.assertEqual(summary["job_score_id"], score["id"])
            self.assertEqual(summary["recommendation"], "review")
            metadata = summary["model_metadata"]
            self.assertEqual(metadata["input_tokens"], 123)
            self.assertEqual(metadata["rubric_reason_code"], "rubric_contract_invalid")
            self.assertEqual(metadata["rubric_rejected_count"], 1)
            self.assertEqual(metadata["rubric_unresolved_requirement_count"], 0)
            questions = journey.supabase.tables["mobile_questions"]
            self.assertTrue(questions)
            self.assertEqual(set(summary["question_ids"]), {row["id"] for row in questions})
            self.assertTrue(all(row["status"] == "pending" and row["job_id"] == job["id"] for row in questions))
            self.assertTrue(any("actual mandatory requirements" in row["prompt"] for row in questions))
            self.assertNotIn("MODEL_ONLY_UNSUPPORTED_CERTIFICATE", json.dumps(journey.supabase.tables))
            self.assertEqual(journey.supabase.tables["artifacts"], [])
            self.assertEqual(provider.operations, ["rank"])
            # Recovery must not loosen the separate candidate-claim contract.
            previous = copy.deepcopy(journey.supabase.tables["job_scores"])
            provider.fabricate_claim = True
            rejected = journey.request("POST", path, json={})
            self.assertEqual(rejected.status_code, 422)
            self.assertEqual(journey.supabase.tables["job_scores"], previous)
            self.assertEqual(journey.supabase.tables["model_runs"][-1]["status"], "failed")
            self.assertNotIn("Managed a hospital team of 50", rejected.text + json.dumps(journey.supabase.tables))
            self.assertEqual(provider.operations, ["rank", "rank"])

    def test_unvalidated_legacy_fallback_cannot_replace_an_existing_score(self):
        with Journey(DeterministicProvider) as journey:
            journey.setup_inputs()
            path = f"jobs/{journey.job['id']}/rank"
            self.assertEqual(journey.request("POST", path, json={}).status_code, 200)
            previous = copy.deepcopy(journey.supabase.tables["job_scores"])
            fallback = studio.StudioResult(
                {"score": 0.0, "recommendation": "review", "rationale": "Unvalidated fallback"},
                model_metadata={"provider": "stub", "model_name": "offline", "status": "received",
                                "input_tokens": 42, "secret": "never-log-this"})
            with patch.object(studio, "rank_job", return_value=fallback):
                self.assertEqual(journey.request("POST", path, json={}).status_code, 502)
            self.assertEqual(journey.supabase.tables["job_scores"], previous)
            audit = journey.supabase.tables["model_runs"][-1]
            self.assertEqual(audit["status"], "failed")
            self.assertEqual(audit["output_summary"]["model_metadata"]["input_tokens"], 42)
            self.assertNotIn("never-log-this", json.dumps(audit))

    def test_nongrounded_document_response_is_rejected_without_artifacts(self):
        class InvalidClaims(DeterministicProvider):
            def complete(self, **kwargs):
                output = json.loads(super().complete(**kwargs))
                output["summary"][0]["source_ids"] = ["profile.experience.99"]
                output["cover_letter"][0]["text"] = "Built unsupported Rust services."
                return json.dumps(output)

        provider = InvalidClaims()
        with Journey(lambda: provider) as journey:
            journey.setup_inputs()
            response = journey.request("POST", f"jobs/{journey.job['id']}/prepare", json={"resume_id": journey.resume["id"]})
            self.assertEqual(response.status_code, 422)
            self.assertEqual(journey.supabase.tables["artifacts"], [])
            self.assertFalse(any(bucket == "application-artifacts" for bucket, _ in journey.supabase.objects))
            self.assertEqual(journey.supabase.tables["model_runs"][0]["status"], "failed")
            self.assertTrue(response.json()["detail"]["questions"])

    def test_cli_live_requires_both_explicit_flags(self):
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        try:
            from mobile_journey_check import main
            for args in (["--execute-live-ai"], ["--confirm-synthetic-data"]):
                with self.subTest(args=args), patch("sys.stdout", new_callable=io.StringIO) as output:
                    self.assertEqual(main(args), 2)
                    self.assertEqual(json.loads(output.getvalue())["status"], "refused")
        finally:
            sys.path.pop(0)

    def test_live_boundary_limits_requests_without_real_network(self):
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        try:
            from mobile_journey_check import LiveBoundary, MAX_LIVE_CALLS
            # Stub only the transport beneath the guard being tested. Never
            # instantiate the credential adapter or contact OpenAI in this test.
            with patch.object(httpx.HTTPTransport, "handle_request", return_value=httpx.Response(200)) as network:
                with LiveBoundary() as boundary, httpx.Client(trust_env=False) as client:
                    for url in ("https://example.test/", "https://api.openai.com/v1/models",
                                "https://api.openai.com.evil.test/v1/chat/completions"):
                        with self.assertRaises(JourneyFailure):
                            client.post(url)
                    self.assertEqual(network.call_count, 0)
                    for _ in range(MAX_LIVE_CALLS):
                        client.post("https://api.openai.com/v1/chat/completions")
                    with self.assertRaises(JourneyFailure):
                        client.post("https://api.openai.com/v1/chat/completions")
                    self.assertEqual(network.call_count, MAX_LIVE_CALLS)
                    self.assertEqual(boundary.http_requests, MAX_LIVE_CALLS)
        finally:
            sys.path.pop(0)

    def test_cli_default_is_hermetic_and_reports_zero_paid_calls(self):
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        try:
            from mobile_journey_check import main
            with patch("sys.stdout", new_callable=io.StringIO) as output:
                self.assertEqual(main([]), 0)
            report = json.loads(output.getvalue())
            self.assertEqual(report["real_openai_http_requests"], 0)
            self.assertEqual(report["live_model_calls"], [])
            self.assertEqual(len(report["artifacts"]), 4)
        finally:
            sys.path.pop(0)

    def test_live_diagnostics_never_echo_model_text_or_source_ids(self):
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        try:
            from mobile_journey_check import output_diagnostics
            raw = json.dumps({"score": 8.0, "recommendation": "match", "questions": [], "requirements": [],
                              "claims": [{"text": "PRIVATE_SYNTHETIC_PROVIDER_TEXT", "source_ids": ["unknown.private.source"]}]})
            result = output_diagnostics(raw, {"operation": "rank", "source_facts": []})
            self.assertFalse(result["claims_valid"])
            self.assertEqual(result["unknown_source_count"], 1)
            self.assertNotIn("PRIVATE_SYNTHETIC", json.dumps(result))
            self.assertNotIn("unknown.private.source", json.dumps(result))
        finally:
            sys.path.pop(0)


if __name__ == "__main__":
    unittest.main()
