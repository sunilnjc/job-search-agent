"""Synthetic multi-profession actual mobile API pipeline evaluation.

Default is offline extraction/persistence. --live uses the previously approved
existing adapter/key for exactly at most ten model requests across five personas.
Auth/DB/Storage remain in-memory; no hosted users, emails, URL fetch, submissions.
"""
import argparse
import base64
import io
import json
import logging
import sys
import time
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'tests'),str(ROOT/'src'),str(ROOT/'scripts')]
from test_mobile_journey import Journey
from mobile_journey_check import LiveBoundary
from jobagent.mobile import studio
from docx import Document
from pypdf import PdfReader

OUT=ROOT/'docs/testing/evidence/pipeline'
LIVE_IDS={'P01','P03','P07','P10','P11'}

def text_of(content,filename):
    if filename.endswith('.pdf'):return '\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(content)).pages)
    return '\n'.join(p.text for p in Document(io.BytesIO(content)).paragraphs)

def evaluate(p,live=False,output_dir=OUT):
    result={'persona':p['id'],'boundary':'real API/extractor/studio; in-memory Supabase','live_ai':live,'checks':{}}
    boundary=LiveBoundary() if live else None
    start=time.monotonic()
    try:
        if boundary:boundary.__enter__()
        with Journey(boundary.provider if live else lambda:None,live=live) as j:
            content=(ROOT/p['resume']).read_bytes(); filename=Path(p['resume']).name
            uploaded=j.request('POST','resumes',json={'filename':filename,'label':p['name'],'content_base64':base64.b64encode(content).decode()})
            result['checks']['upload_status']=uploaded.status_code
            if uploaded.status_code!=201:return result
            rid=uploaded.json()['id']
            extracted=j.request('GET',f'resumes/{rid}/text')
            result['checks']['extract_status']=extracted.status_code
            extracted_text=extracted.json().get('text','')
            norm=lambda t:' '.join(t.split())
            result['checks']['facts_found']=[line for line in p['facts'] if norm(line) in norm(extracted_text)]
            result['checks']['facts_missing']=[line for line in p['facts'] if norm(line) not in norm(extracted_text)]
            result['checks']['name_found']=p['name'] in extracted_text
            profile={'display_name':p['name'],'phone':p['phone'],'base_location':p['location'],'career_text':'\n'.join(p['facts']),'career_background':{'profession':p['profession'],'experience_level':p['career_stage'],'qualifications':[]}}
            saved=j.request('PUT','profile',json=profile);result['checks']['profile_status']=saved.status_code
            prefs=j.request('PUT','preferences',json={'target_titles':p['roles'],'preferred_locations':p['locations'],'remote_preference':p['remote'],'sponsorship_required':p['sponsorship'],'work_authorization_notes':p['constraints']})
            result['checks']['preferences_status']=prefs.status_code
            job=j.request('POST','jobs',json=p['job']);result['checks']['job_status']=job.status_code
            jid=job.json()['id']
            result['checks']['cross_user_download_status']=j.request('GET',f'resumes/{rid}/download',actor='b').status_code
            if live:
                rank=j.request('POST',f'jobs/{jid}/rank',json={'resume_id':rid})
                result['rank']={'status':rank.status_code,'body':rank.json()}
                prepared=j.request('POST',f'jobs/{jid}/prepare',json={'resume_id':rid,'variant':'career_change' if p['career_stage']=='career_change' else 'role_aligned'})
                result['prepare']={'status':prepared.status_code,'artifacts':[]}
                if prepared.status_code!=200:result['prepare']['error']=prepared.json()
                for artifact in prepared.json().get('artifacts',[]):
                    response=j.request('GET',f"artifacts/{artifact['id']}/download")
                    path=output_dir/(p['id']+'-'+artifact['kind']+Path(artifact['filename']).suffix)
                    path.write_bytes(response.content)
                    text=text_of(response.content,path.name)
                    result['prepare']['artifacts'].append({'kind':artifact['kind'],'format':path.suffix,'download_status':response.status_code,'path':str(path.relative_to(ROOT)),'text':text,'source_ids':artifact.get('source_ids',[])})
                result['model_calls']=boundary.calls
            result['checks']['no_applications_submitted']=not any(x.get('status')=='submitted' for x in j.supabase.tables.get('applications',[]))
            result['checks']['source_download_identical']=j.request('GET',f'resumes/{rid}/download').content==content
    except Exception as error:
        result['error_type']=type(error).__name__ # Never print raw SDK payload/errors.
    finally:
        if boundary:boundary.__exit__(None,None,None)
        result['elapsed_seconds']=round(time.monotonic()-start,2)
    return result

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--live',action='store_true')
    parser.add_argument('--output-dir', type=Path, default=OUT)
    parser.add_argument('--persona', action='append', choices=sorted(LIVE_IDS))
    args=parser.parse_args()
    logging.disable(logging.CRITICAL)
    output_dir=args.output_dir.resolve()
    if not output_dir.is_relative_to(ROOT/'docs/testing/evidence'):
        parser.error('Use a folder under docs/testing/evidence; no external destination is allowed.')
    if output_dir.exists() and any(output_dir.iterdir()):
        parser.error('Choose a new evidence folder to preserve previous test results.')
    output_dir.mkdir(parents=True,exist_ok=True)
    results=[]
    for p in json.loads((ROOT/'tests/fixtures/resumes/personas.json').read_text()):
        if args.persona and p['id'] not in args.persona:continue
        result=evaluate(p,live=args.live and p['id'] in LIVE_IDS,output_dir=output_dir);results.append(result)
        (output_dir/'results.json').write_text(json.dumps(results,indent=2,ensure_ascii=False)+'\n')
        print(p['id'],json.dumps({k:v for k,v in result['checks'].items() if k not in {'facts_found','facts_missing'}}),'missing facts',len(result['checks'].get('facts_missing',[])),'AI',result.get('prepare',{}).get('status','not called'),flush=True)
    return 0 if all(not r.get('error_type') and (not r['live_ai'] or (r.get('rank',{}).get('status')==200 and r.get('prepare',{}).get('status')==200)) for r in results) else 1

if __name__=='__main__':raise SystemExit(main())
