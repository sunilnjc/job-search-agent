"""Single-column, text-first exports. No network, templates, or candidate defaults.

The studio supplies validated text and confirmed contact links. These renderers
escape all text, share a conservative page plan, and reject overflow rather than
shrinking type or silently dropping experience. Layout bytes are deterministic;
the caller assigns unique, versioned filenames.
"""
from __future__ import annotations

import io
import ipaddress
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit
from xml.sax.saxutils import escape, quoteattr
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
import reportlab
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate


@dataclass(frozen=True)
class Block:
    text: str
    style: str = "body"
    href: str = ""


def clean_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = value.translate(str.maketrans({"–": "-", "—": "-", "‑": "-", "−": "-"}))
    return " ".join("".join(c for c in value if c in "\n\t" or
                            not unicodedata.category(c).startswith("C")).split())


def safe_link(value: str) -> str:
    """Only explicit public web URLs or bare mailto addresses; never fetch them."""
    if not isinstance(value, str) or len(value) > 500 or re.search(r"[\s<>\"'\\]", value):
        return ""
    if any(unicodedata.category(c).startswith("C") for c in value):
        return ""
    if value.startswith("mailto:"):
        return value if re.fullmatch(r"mailto:[A-Za-z0-9.!#$%&*+/=_`{|}~-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", value) else ""
    try:
        parts = urlsplit(value)
        host = parts.hostname or ""
        if parts.scheme not in {"https", "http"} or not host or parts.username or parts.password:
            return ""
        if parts.port not in {None, 80, 443} or "." not in host or host.endswith((".local", ".internal")):
            return ""
        try:
            ipaddress.ip_address(host)
            return ""  # Contact profiles do not need IP-literal links.
        except ValueError:
            pass
        if not re.fullmatch(r"[A-Za-z0-9.-]+", host) or "%" in value:
            return ""
        return value
    except ValueError:
        return ""


def _fonts() -> tuple[str, str]:
    # Vera ships with our existing reportlab dependency; no host-specific fonts.
    directory = Path(reportlab.__file__).parent / "fonts"
    for name, filename in (("MobileVera", "Vera.ttf"), ("MobileVeraBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(directory / filename)))
    return "MobileVera", "MobileVeraBold"


def _styles() -> dict:
    regular, bold = _fonts()
    return {
        "title": ParagraphStyle("title", fontName=bold, fontSize=18, leading=23, spaceAfter=8),
        "heading": ParagraphStyle("heading", fontName=bold, fontSize=12, leading=16,
                                  spaceBefore=9, spaceAfter=5, keepWithNext=True),
        "body": ParagraphStyle("body", fontName=regular, fontSize=11, leading=15, spaceAfter=6),
        "contact": ParagraphStyle("contact", fontName=regular, fontSize=10, leading=14, spaceAfter=4),
    }


def _paragraph(block: Block, styles: dict) -> Paragraph:
    text = escape(block.text)
    if block.href:
        text = f"<link href={quoteattr(block.href)} color='black'>{text}</link>"
    return Paragraph(text, styles[block.style])


def _pages(blocks: list[Block], max_pages: int) -> list[list[Block]]:
    styles = _styles()
    cleaned = [Block(clean_text(b.text), b.style, b.href) for b in blocks if clean_text(b.text)]
    glyphs = pdfmetrics.getFont("MobileVera").face.charToGlyph
    if not cleaned:
        raise ValueError("Add confirmed career details before exporting.")
    if any(ord(c) not in glyphs for b in cleaned for c in b.text):
        raise ValueError("This export font cannot display some characters. Use Latin-script text for now.")
    if any(b.style not in styles or (b.href and safe_link(b.href) != b.href) for b in cleaned):
        raise ValueError("The document contains an unsupported style or contact link.")
    # Letter, 0.75-inch margins. Reserve extra width/height for Word font metrics.
    width, height = 474, 650
    heights = [_paragraph(b, styles).wrap(width, height)[1] + styles[b.style].spaceBefore
               + styles[b.style].spaceAfter for b in cleaned]
    pages: list[list[Block]] = [[]]
    used = 0.0
    for index, block in enumerate(cleaned):
        needed = heights[index]
        if block.style == "heading" and index + 1 < len(cleaned):
            needed += heights[index + 1]
        if needed > height:
            raise ValueError("A paragraph is too long. Shorten it before exporting.")
        if used + needed > height:
            pages.append([])
            used = 0
        pages[-1].append(block)
        used += heights[index]
    if len(pages) > max_pages:
        raise ValueError(f"This draft exceeds {max_pages} pages. Select fewer confirmed details and try again.")
    return pages


def render_pdf(blocks: list[Block], *, max_pages: int = 2) -> bytes:
    pages = _pages(blocks, max_pages)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=(612, 792), leftMargin=54, rightMargin=54,
                            topMargin=54, bottomMargin=54, title="", author="", invariant=1)
    styles = _styles()
    story = []
    for index, page in enumerate(pages):
        if index:
            story.append(PageBreak())
        story.extend(_paragraph(b, styles) for b in page)
    doc.build(story)
    result = buffer.getvalue()
    from pypdf import PdfReader
    if len(PdfReader(io.BytesIO(result)).pages) > max_pages:
        raise ValueError("The draft overflowed its page budget. Shorten the draft and retry.")
    return result


def render_docx(blocks: list[Block], *, max_pages: int = 2) -> bytes:
    pages = _pages(blocks, max_pages)
    document = Document()
    # Some python-docx distributions ship a decorated Title style. Do not inherit
    # theme rules/borders in a plain, single-column ATS export.
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.75)
    for name, size in (("Normal", 11), ("Title", 18), ("Heading 1", 12)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != "Normal"
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = Pt(15 if name == "Normal" else size + 5)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(9)
    for page_index, page in enumerate(pages):
        if page_index:
            document.add_page_break()
        for block in page:
            style = {"title": "Title", "heading": "Heading 1"}.get(block.style, "Normal")
            paragraph = document.add_paragraph(style=style)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = block.style in {"title", "heading"}
            if block.style == "contact":
                paragraph.paragraph_format.line_spacing = Pt(14)
                paragraph.paragraph_format.space_after = Pt(4)
            if block.href:
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("r:id"), paragraph.part.relate_to(block.href, RT.HYPERLINK, is_external=True))
                run = OxmlElement("w:r")
                text = OxmlElement("w:t")
                text.text = block.text
                run.append(text)
                hyperlink.append(run)
                paragraph._p.append(hyperlink)
            else:
                run = paragraph.add_run(block.text)
                if block.style == "contact":
                    run.font.size = Pt(10)
    properties = document.core_properties
    properties.author = properties.last_modified_by = properties.title = properties.subject = ""
    properties.comments = properties.keywords = ""
    properties.created = properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    raw = io.BytesIO()
    document.save(raw)
    # python-docx uses wall-clock ZIP metadata; normalize it for reproducible bytes.
    output = io.BytesIO()
    with ZipFile(raw) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            entry = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            target.writestr(entry, source.read(name))
    return output.getvalue()
